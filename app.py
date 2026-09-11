from flask import Flask, render_template, request, jsonify, send_file
from detector import detect_sensitive_data
from docx import Document
from pypdf import PdfReader
import io
import re

app = Flask(__name__)


def extract_text(file):
    filename = file.filename.lower()

    # TXT files
    if filename.endswith(".txt"):
        return file.read().decode("utf-8", errors="ignore")

    # DOCX files
    elif filename.endswith(".docx"):
        document = Document(file)
        text = []

        # Read normal paragraphs
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)

        # Read table contents
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)

        return "\n".join(text)

    # PDF files
    elif filename.endswith(".pdf"):
        reader = PdfReader(file)
        text = []

        for page in reader.pages:
            page_text = page.extract_text()

            if page_text:
                text.append(page_text)

        return "\n".join(text)

    return ""


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/scan", methods=["POST"])
def scan():

    if "file" not in request.files:
        return jsonify({
            "findings": [],
            "count": 0,
            "score": 0,
            "risk_level": "SAFE"
        })

    file = request.files["file"]

    if file.filename == "":
        return jsonify({
            "findings": [],
            "count": 0,
            "score": 0,
            "risk_level": "SAFE"
        })

    text = extract_text(file)

    print("EXTRACTED TEXT:")
    print(text)

    findings, score, risk_level = detect_sensitive_data(text)

    return jsonify({
        "findings": findings,
        "count": len(findings),
        "score": score,
        "risk_level": risk_level
    })


@app.route("/make-safe", methods=["POST"])
def make_safe():

    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    file = request.files["file"]

    text = extract_text(file)

    # Mask Email
    text = re.sub(
        r'\b([A-Za-z0-9._%+-])[A-Za-z0-9._%+-]*(@[A-Za-z0-9.-]+\.[A-Za-z]{2,})\b',
        r'\1******\2',
        text
    )

    # Mask Phone
    text = re.sub(
        r'\b([6-9])\d{7}(\d{2})\b',
        r'\1*******\2',
        text
    )

    # Mask Aadhaar-like number
    text = re.sub(
        r'\b\d{4}\s?\d{4}\s?\d{4}\b',
        'XXXX XXXX XXXX',
        text
    )

    # Mask PAN-like number
    text = re.sub(
        r'\b[A-Z]{5}[0-9]{4}[A-Z]\b',
        'XXXXXXXXXX',
        text
    )

    # Mask Account Number
    text = re.sub(
        r'(?i)\b(account|a/c|acct)[\s#:_-]*\d{8,18}\b',
        r'\1: [REDACTED]',
        text
    )

    # Mask API keys / secrets / tokens
    text = re.sub(
        r'(?i)(api[_-]?key|secret|token)\s*[:=]\s*[A-Za-z0-9_\-]{8,}',
        r'\1=[REDACTED]',
        text
    )

    safe_file = io.BytesIO(text.encode("utf-8"))
    safe_file.seek(0)

    return send_file(
        safe_file,
        as_attachment=True,
        download_name="safe_copy.txt",
        mimetype="text/plain"
    )


if __name__ == "__main__":
    app.run(debug=True)